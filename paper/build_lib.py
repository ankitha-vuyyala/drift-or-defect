"""Small helpers for building the paper as .docx with python-docx (numbered headings, captions, tables, figures,
IEEE-style numbered references, and an evidence index that maps every reported number to its source)."""
from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Times New Roman"


class Paper:
    def __init__(self):
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)        # US Letter
        for side in ("left_margin", "right_margin"):
            setattr(sec, side, Cm(2.2))
        sec.top_margin = sec.bottom_margin = Cm(2.2)
        st = self.doc.styles["Normal"]
        st.font.name, st.font.size = FONT, Pt(10)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.08
        for lvl, size in ((1, 12), (2, 10.5), (3, 10)):
            h = self.doc.styles[f"Heading {lvl}"]
            h.font.name, h.font.size, h.font.bold = FONT, Pt(size), True
            h.font.color.rgb = RGBColor(0, 0, 0)
            h.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
            h.paragraph_format.space_before = Pt(10 if lvl == 1 else 6)
            h.paragraph_format.space_after = Pt(3)
        self.sec = [0, 0, 0]
        self.fig_n = self.tab_n = 0
        self.refs: list[str] = []          # keys in citation order
        self.ref_db: dict[str, str] = {}
        self.evidence: list[dict] = []
        self.labels: dict[str, str] = {}

    # ------------------------------------------------------------------ text
    def title(self, text: str, authors: str, note: str = ""):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold, r.font.size = True, Pt(16)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(authors).italic = True
        if note:
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(note)
            r.font.size = Pt(8.5)

    def heading(self, text: str, level: int = 1, numbered: bool = True, label: str | None = None,
                anchor: str | None = None):
        if not numbered and label and anchor:
            self.labels[label] = anchor
        if numbered:
            self.sec[level - 1] += 1
            for i in range(level, 3):
                self.sec[i] = 0
            num = ".".join(str(x) for x in self.sec[:level])
            if label:
                self.labels[label] = num
            text = f"{num} {text}"
        self.doc.add_heading(text, level=level)

    def para(self, text: str, style: str | None = None, italic: bool = False, size: float | None = None,
             align_justify: bool = True):
        """Paragraph with **bold**, *italic* and [@ref] citations, {fig:x}/{tab:x}/{sec:x} cross references."""
        p = self.doc.add_paragraph(style=style)
        if align_justify:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        text = self._cites(self._xrefs(text))
        for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
            if not part:
                continue
            if part.startswith("**"):
                r = p.add_run(part[2:-2])
                r.bold = True
            elif part.startswith("*"):
                r = p.add_run(part[1:-1])
                r.italic = True
            else:
                r = p.add_run(part)
            if italic:
                r.italic = True
            if size:
                r.font.size = Pt(size)
        return p

    def bullets(self, items: list[str]):
        for it in items:
            p = self.para(it, style="List Bullet")
            p.paragraph_format.space_after = Pt(1)

    def equation(self, text: str):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.italic = True

    def _xrefs(self, text: str) -> str:
        def rep(m):
            kind, key = m.group(1), m.group(2)
            v = self.labels.get(f"{kind}:{key}")
            if kind == "sec" and v and v.startswith("Appendix"):
                return v
            return {"fig": "Fig. ", "tab": "Table ", "sec": "Section "}[kind] + (v or "??")
        return re.sub(r"\{(fig|tab|sec):([A-Za-z0-9_\-]+)\}", rep, text)

    def _cites(self, text: str) -> str:
        def rep(m):
            keys = [k.strip().lstrip("@") for k in m.group(1).split(";")]
            nums = []
            for k in keys:
                if k not in self.ref_db:
                    raise KeyError(f"unknown reference {k}")
                if k not in self.refs:
                    self.refs.append(k)
                nums.append(self.refs.index(k) + 1)
            return "[" + ", ".join(str(n) for n in sorted(nums)) + "]"
        return re.sub(r"\[@([^\]]+)\]", rep, text)

    # ------------------------------------------------------------------ floats
    def figure(self, path: Path, caption: str, label: str, width_cm: float = 16.0):
        self.fig_n += 1
        self.labels[f"fig:{label}"] = str(self.fig_n)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Cm(width_cm))
        c = self.doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = c.add_run(f"Fig. {self.fig_n}. ")
        r.bold = True
        r.font.size = Pt(9)
        r2 = c.add_run(self._cites(self._xrefs(caption)))
        r2.font.size = Pt(9)

    def table(self, header: list[str], rows: list[list[str]], caption: str, label: str,
              col_widths_cm: list[float] | None = None, font_pt: float = 8.0, bold_rows: set | None = None):
        self.tab_n += 1
        self.labels[f"tab:{label}"] = str(self.tab_n)
        c = self.doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = c.add_run(f"Table {self.tab_n}. ")
        r.bold = True
        r.font.size = Pt(9)
        r2 = c.add_run(self._cites(self._xrefs(caption)))
        r2.font.size = Pt(9)
        t = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.style = "Table Grid"
        for j, h in enumerate(header):
            cell = t.rows[0].cells[j]
            cell.text = ""
            run = cell.paragraphs[0].add_run(h)
            run.bold = True
            run.font.size = Pt(font_pt)
            _shade(cell, "E8E8E8")
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                cell = t.rows[i + 1].cells[j]
                cell.text = ""
                run = cell.paragraphs[0].add_run(str(v))
                run.font.size = Pt(font_pt)
                if bold_rows and i in bold_rows:
                    run.bold = True
        if col_widths_cm:
            t.autofit = False
            tblPr = t._tbl.tblPr
            lay = OxmlElement("w:tblLayout")
            lay.set(qn("w:type"), "fixed")
            tblPr.append(lay)
            for j, w in enumerate(col_widths_cm):
                t.columns[j].width = Cm(w)
            grid = t._tbl.tblGrid
            for j, gc in enumerate(grid.findall(qn("w:gridCol"))):
                gc.set(qn("w:w"), str(int(col_widths_cm[j] * 567)))
            for row in t.rows:
                for j, w in enumerate(col_widths_cm):
                    row.cells[j].width = Cm(w)
        self.doc.add_paragraph()

    # ------------------------------------------------------------------ references and evidence
    def load_refs(self, path: Path):
        self.ref_db = json.loads(path.read_text(encoding="utf-8"))

    def references(self):
        self.heading("References", 1, numbered=False)
        for i, k in enumerate(self.refs, 1):
            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.7)
            p.paragraph_format.first_line_indent = Cm(-0.7)
            r = p.add_run(f"[{i}] {self.ref_db[k]}")
            r.font.size = Pt(8.5)

    def ev(self, value, fmt: str, source: str, what: str) -> str:
        """Format a reported number and record where it comes from (evidence index)."""
        s = format(value, fmt) if not isinstance(value, str) else value
        self.evidence.append({"value": s, "source": source, "what": what})
        return s

    def save(self, path: Path):
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(path))


def _shade(cell, hex_fill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)
